#!/usr/bin/env python3
"""One-time migration: merge the source spreadsheets into a single master workbook.

  tools/explain/.venv/bin/python tools/explain/export_master.py "<out.xlsx>" [src1.xlsx src2.xlsx ...]

Sources default to tools/sources.local.txt. Rows are copied raw (so question ids, which hash
the text, do not change), duplicates keep their first occurrence, a row shifted one column
left is repaired, and the explanation pipeline's results fill the review columns:

  K  Answer explanation    published   (ready explanations)
  L  Exhibit               published   (empty until you approve one)
  M  Review status         not published: why the row needs the instructor
  N  Proposed explanation  not published: copy into K to approve
  O  Proposed exhibit      not published: copy into L to approve
"""
import hashlib, os, re, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))
from build_questions import TYPES, clean                     # noqa: E402
import check                                                 # noqa: E402
from common import LETTERS, bank, letters, read              # noqa: E402

HEADER = ['Question Text', 'Question Type', 'Option 1', 'Option 2', 'Option 3', 'Option 4', 'Option 5',
          'Correct Answer', 'Time in seconds', 'Image Link', 'Answer explanation', 'Exhibit',
          'Review status', 'Proposed explanation', 'Proposed exhibit']
WIDTHS = {'A': 80, 'B': 16, 'C': 28, 'D': 28, 'E': 28, 'F': 28, 'G': 28, 'H': 10, 'I': 9, 'J': 12,
          'K': 60, 'L': 60, 'M': 50, 'N': 60, 'O': 60}
MONO = {'L', 'O'}
WRAP = {'A', 'C', 'D', 'E', 'F', 'G', 'K', 'L', 'M', 'N', 'O'}

def source_rows(path):
    ws = openpyxl.load_workbook(path).active
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 12)]       # A..K, raw
        if not any(v not in (None, '') for v in vals):
            continue
        b = str(vals[1] or '').strip()
        if b and b.lower() not in TYPES:                           # cells slid one column left
            vals = [vals[0], 'Multiple Choice'] + vals[1:10]
        yield r, vals

def nums(lets):
    return ','.join(str(LETTERS.index(x) + 1) for x in lets)

def review(q, r1, r2, ex, v):
    """-> (K, M, N, O) for this question from the pipeline results."""
    status, expl, exhibit, reasons = check.assess(q, r1, r2, ex, v)
    why = '; '.join(reasons)
    if status == 'ready':
        return expl, '', '', ''
    if status == 'NEW EXHIBIT':
        return '', ('NEW EXHIBIT: generated to replace the one lost in the Excel export; a blind check '
                    'picked the keyed answer. To publish, copy O into L and N into K.'), expl, exhibit
    if status == 'KEY WRONG?':
        better = r2['better_answer']
        m = (f"KEY WRONG? Suggested answer: {nums(better)} "
             f"({' + '.join(q['options'][LETTERS.index(x)] for x in better)}). {r2['note']}")
        n = ''
        if r1 and sorted(r1['answer']) == sorted(better):
            n = r1['explanation']
            m += ' Column N explains the suggested answer; if you change H, copy N into K.'
        return '', m, n, ''
    if status == 'EXHIBIT: drop':
        return '', f'MISSING EXHIBIT: could not be rebuilt ({why}). Delete this row or add an exhibit in L.', '', exhibit
    if status == 'EXHIBIT: pending':
        return '', f'MISSING EXHIBIT: not processed yet ({why}).', '', ''
    note = (r2 or {}).get('note') or (r1 or {}).get('note') or ''
    return '', f"{status.upper()}: {why}." + (f' {note}' if note else '') + (
        ' If N is right, copy it into K.' if expl else ''), expl, ''

def main():
    out = sys.argv[1]
    srcs = sys.argv[2:] or [l.strip() for l in open(os.path.join(os.path.dirname(HERE), 'sources.local.txt'))
                            if l.strip() and not l.startswith('#')]
    B = {q['id']: q for q in bank()}
    p1, p2, ex, ver = read('pass1.json'), read('pass2.json'), read('exhibits.json'), read('verify.json')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Questions'
    ws.append(HEADER)
    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(name='Arial', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='FFD8D8D8')
    seen, counts = set(), {'rows': 0, 'dupes': 0, 'K': 0, 'M': 0}
    for src in srcs:
        for r, vals in source_rows(src):
            qid = hashlib.sha1(clean(str(vals[0])).encode('utf-8')).hexdigest()[:10]
            if qid in seen:
                counts['dupes'] += 1
                continue
            seen.add(qid)
            q = B.get(qid)
            k, m, n, o = review(q, p1.get(qid), p2.get(qid), ex.get(qid), ver.get(qid)) if q else ('', '', '', '')
            if not q:
                m = 'NOT IN BANK: the converter skips this row (missing answer or options).'
            ws.append(vals[:10] + [k or vals[10], '', m, n, o])
            counts['rows'] += 1
            counts['K'] += bool(k or vals[10])
            counts['M'] += bool(m)

    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    body, mono = Font(name='Aptos Narrow', size=11), Font(name='Consolas', size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            col = get_column_letter(cell.column)
            cell.font = mono if col in MONO else body
            cell.alignment = Alignment(wrap_text=col in WRAP, vertical='top')
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:O{ws.max_row}'
    wb.save(out)
    print(f"wrote {out}: {counts['rows']} questions ({counts['dupes']} duplicates dropped), "
          f"{counts['K']} with an explanation, {counts['M']} needing review")

if __name__ == '__main__':
    main()
